#!/usr/bin/env python3
"""Run mudshark_summary.py on the fixture and diff against the expected file.

- Reads BOTH files with the tool's own stdlib reader (consistency).
- ALSO opens the generated file with openpyxl to prove the writer produces
  a valid, Excel-compatible .xlsx.
"""
import os
import subprocess
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)

from mudshark_summary import read_workbook  # noqa: E402

FIXTURE = os.path.join(HERE, "fixture_result.xlsx")
EXPECTED = os.path.join(HERE, "expected_summary.xlsx")
ACTUAL = os.path.join(HERE, "actual_summary.xlsx")


def sheet_matrix(path):
    book = read_workbook(path)
    rows = next(iter(book.values()))  # first (and only) sheet
    return {rn: cells for rn, _hidden, cells in rows}


def nearly(a, b):
    if isinstance(a, (int, float)) and isinstance(b, (int, float)) \
            and not isinstance(a, bool) and not isinstance(b, bool):
        return abs(float(a) - float(b)) <= 1e-9 * max(1.0, abs(float(a)))
    return a == b


def compare():
    exp = sheet_matrix(EXPECTED)
    act = sheet_matrix(ACTUAL)
    problems = 0
    for i in sorted(set(exp) | set(act)):
        e = exp.get(i, {})
        a = act.get(i, {})
        cols = sorted(set(e) | set(a))
        for c in cols:
            ev = e.get(c)
            av = a.get(c)
            if ev in (" ", ""):
                ev = None
            if av in (" ", ""):
                av = None
            if ev is None and av is None:
                continue
            if isinstance(ev, str):
                ev = ev.strip()
            if isinstance(av, str):
                av = av.strip()
            if not nearly(ev, av):
                print("  MISMATCH row %-3d col %-2d expected=%r actual=%r"
                      % (i, c, ev, av))
                problems += 1
    return problems


def main():
    print("== running mudshark_summary.py on fixture ==")
    r = subprocess.run([sys.executable,
                        os.path.join(ROOT, "mudshark_summary.py"),
                        FIXTURE, "-o", ACTUAL],
                       capture_output=True, text=True)
    print(r.stdout)
    if r.returncode != 0:
        print(r.stderr)
        return 1

    print("== diff actual vs expected ==")
    problems = compare()
    if problems:
        print("FAILED: %d mismatch(es)" % problems)
        return 1
    print("OK - generated Summary matches the hand-made one cell-for-cell.")

    print("== validating output with openpyxl ==")
    try:
        import openpyxl
        wb = openpyxl.load_workbook(ACTUAL)
        ws = wb.active
        print("openpyxl opened '%s' fine: sheet=%r size=%dx%d"
              % (os.path.basename(ACTUAL), ws.title,
                 ws.max_row, ws.max_column))
    except ImportError:
        print("(openpyxl not available here - skipped)")
    except Exception as exc:  # noqa: BLE001
        print("FAILED: openpyxl could not open the generated file:", exc)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
